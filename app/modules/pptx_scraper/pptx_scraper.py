"""pptx_scraper.py

_summary_

_extended_summary_

Returns:
    _type_: _description_

Ex:
pptx_path = "path_to_your_pptx_file.pptx"
csv_path = "output_file.csv"
scraper = PPTXScraper(pptx_path)
scraper.save_to_csv(csv_path)
"""

import csv
from pptx import Presentation


class PPTXScraper:
    """
     _summary_

    _extended_summary_
    """

    def __init__(self, pptx_path):
        """
        __init__ _summary_

        _extended_summary_

        Args:
            pptx_path (_type_): _description_
        """
        self.pptx_path = pptx_path
        self.presentation = Presentation(pptx_path)

    def _get_slide_text(self, slide):
        """
        _get_slide_text _summary_

        _extended_summary_

        Args:
            slide (_type_): _description_

        Returns:
            _type_: _description_
        """
        text = ""
        for shape in slide.shapes:
            if hasattr(shape, "text"):
                text += shape.text + " "
        return text.strip()

    def _get_slide_notes(self, slide):
        """
        _get_slide_notes _summary_

        _extended_summary_

        Args:
            slide (_type_): _description_

        Returns:
            _type_: _description_
        """
        if slide.has_notes_slide:
            return slide.notes_slide.notes_text_frame.text
        return ""

    def scrape(self):
        """
        scrape _summary_

        _extended_summary_

        Returns:
            _type_: _description_
        """
        data = []
        for i, slide in enumerate(self.presentation.slides, start=1):
            slide_text = self._get_slide_text(slide)
            notes_text = self._get_slide_notes(slide)
            data.append([i, slide_text, notes_text])
        return data

    def save_to_csv(self, csv_path):
        """
        save_to_csv _summary_

        _extended_summary_

        Args:
            csv_path (_type_): _description_
        """
        data = self.scrape()
        with open(csv_path, "w", newline="", encoding="utf-8") as file:
            writer = csv.writer(file)
            writer.writerow(["Slide Number", "Slide Text", "Notes"])
            writer.writerows(data)
